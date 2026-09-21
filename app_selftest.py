"""Opt-in executable integration check; inputs are provided on the command line."""
import json
from pathlib import Path
import time
from unittest.mock import patch

from openpyxl import load_workbook
from battery_schema import profile, available_columns, display_values
from neware_batch import json_write, DEFAULT_COLUMNS


def start_without_prompt(panel):
    from tkinter import filedialog
    choose_dir, choose_file = filedialog.askdirectory, filedialog.asksaveasfilename
    def unexpected(**_):
        raise AssertionError('Saved output location must not trigger another dialog')
    filedialog.askdirectory = filedialog.asksaveasfilename = unexpected
    try:
        panel.start_button.invoke()
        assert panel.busy
    finally:
        filedialog.askdirectory, filedialog.asksaveasfilename = choose_dir, choose_file


def wait_for(root, panel):
    deadline = time.monotonic() + 180
    while panel.busy and time.monotonic() < deadline:
        root.update(); time.sleep(.02)
    assert not panel.busy, 'Background work did not finish'


def run_selftest(app, args):
    assert args.out and args.batch
    args.out = args.out.resolve()
    root = app.root
    callback_errors = []
    root.report_callback_exception = lambda kind, value, tb: callback_errors.append(str(value))
    args.out.mkdir(parents=True, exist_ok=True)
    from drop_selftest import check_drop_inputs, native_drop
    app.root.update_idletasks()
    assert not app.files and not app.text_panel.files
    drop_checks = check_drop_inputs(app)
    assert len(app.notebook.tabs()) == 3
    for index, button in enumerate(app.nav_buttons):
        button.invoke(); root.update_idletasks()
        assert app.notebook.index('current') == index
    app.nav_buttons[0].invoke()
    app.metrics_button.invoke(); assert app.extra_fields_open
    app.metrics_button.invoke(); assert not app.extra_fields_open
    assert not hasattr(app, 'preview') and not hasattr(app.text_panel, 'preview')
    for tab in app.notebook.tabs(): app.notebook.select(tab); root.update_idletasks()
    app.notebook.select(0)
    assert native_drop(app, args.batch) == 'copy'
    initial = tuple(range(8)) if args.all_columns else DEFAULT_COLUMNS
    assert app.current_columns() == initial
    start_without_prompt(app)
    wait_for(root, app)
    manifest = app.last_manifest
    assert manifest and manifest['failed'] == 0 and not manifest['excel_error'], manifest
    assert manifest['statistics']['success_rate_percent'] == 100
    assert '提取成功率 100.0%' in app.status.get()
    checks = []
    workbook = load_workbook(manifest['workbook'])
    for entry in manifest['entries']:
        app.tree.selection_set(str(entry['index'])); app.show_selection()
        result = json.loads(Path(entry['data']).read_text(encoding='utf-8')); schema = profile(result)
        selected = available_columns(initial, schema)
        assert app.tree.item(str(entry['index']))['values'][0] == Path(entry['source']).name
        sheet = workbook[entry['sheet_name']]
        assert list(next(sheet.values)) == [schema['headers'][i] for i in selected]
        assert sheet.max_row == entry['cycles'] + 1
        for row, cells in zip(result['cycles'], sheet.iter_rows(min_row=2)):
            expected = display_values(row, schema)
            assert [c.value for c in cells] == [float(expected[i]) if expected[i] else None for i in selected]
            assert all(c.value is None or c.data_type == 'n' for c in cells)
        checks.append({'name': entry['sample_name'], 'rows': entry['cycles'], 'columns': len(selected)})
    workbook.close()
    for i, v in app.field_vars.items():
        if not v.get(): app.field_buttons[i].invoke()
    target = args.out / '另存测试' / '循环数据_全部列.xlsx'
    from tkinter import filedialog
    choose_name = filedialog.asksaveasfilename
    filedialog.asksaveasfilename = lambda **_: str(target)
    try:
        app.actions_menu.invoke('另存所选列')
    finally:
        filedialog.asksaveasfilename = choose_name
    assert target.exists()
    workbook = load_workbook(target)
    for sheet, entry in zip(workbook, manifest['entries']):
        schema = profile(json.loads(Path(entry['data']).read_text(encoding='utf-8')))
        assert list(next(sheet.values)) == schema['headers']
    workbook.close()
    for button in app.field_buttons: button.invoke()
    assert not app.current_columns() and str(app.actions_menu.entrycget('另存所选列', 'state')) == 'disabled'
    for i in DEFAULT_COLUMNS: app.field_buttons[i].invoke()
    from land_selftest import check_statistics_choice
    policy_checks = check_statistics_choice(app, manifest, args.out, wait_for, start_without_prompt)
    text_manifest = None
    feedback_checks = {}
    if args.text_folder:
        panel = app.text_panel; app.notebook.select(1); root.update_idletasks()
        assert panel.cv_split.get() and not panel.eis_simple.get()
        panel.input.set(str(args.text_folder)); panel.output.set(str(args.out / '文本界面测试'))
        panel.scan_folder()
        assert len(panel.tree.get_children()) == 3
        assert panel.advanced_window.state() == 'withdrawn'
        panel.toggle_settings(); root.update_idletasks()
        # A transient dialog follows the deliberately hidden smoke-test root.
        if root.state() != 'withdrawn': assert panel.advanced_window.state() == 'normal'
        panel.advanced_window.withdraw()
        panel.eis_simple.set(True); panel.cv_split.set(False)
        panel.low.set('1.5'); panel.high.set('2.2'); panel.add_range()
        assert panel.ranges == [(1.5, 2.2, 'min')]
        start_without_prompt(panel)
        wait_for(root, panel)
        text_manifest = panel.manifest
        assert text_manifest and not text_manifest['failed'] and not text_manifest['excel_errors'], text_manifest
        assert text_manifest['statistics']['success_rate_percent'] == 100
        assert '提取成功率 100.0%' in panel.status.get()
        assert len(text_manifest['workbooks']) == 4
        for entry in text_manifest['entries']:
            assert Path(entry['workbook']).is_relative_to(args.out)
            panel.tree.selection_set(str(entry['index'])); panel.show_selection()
            assert entry['workbook'] in panel.detail.get()
            book = load_workbook(entry['workbook'], read_only=True)
            sheet = book[entry['sheet_name']]
            if entry['kind'] == 'EIS': assert sheet.max_column == 2
            if entry['kind'] == 'CV':
                source = json.loads(Path(entry['data']).read_text(encoding='utf-8'))
                assert sheet.max_column == len(source['headers']) and sheet.max_row == entry['rows'] + 1
            book.close()
        cases = args.out/'失败流程输入'; cases.mkdir(exist_ok=True)
        good, bad = cases/'good.txt', cases/'bad.txt'
        good.write_text('x,y\n1,2\n3,4', encoding='utf-8'); bad.write_bytes(b'\xff\xfe\x00')
        panel.clear_files(); panel.add_paths([good, bad]); panel.output.set(str(args.out/'文本失败流程'))
        with patch('failed_files.messagebox.askyesno', return_value=False) as ask:
            start_without_prompt(panel); wait_for(root, panel); ask.assert_called_once()
        assert panel.manifest['statistics']['success_rate_percent'] == 50
        assert '提取成功率 50.0%' in panel.status.get()
        assert not list(Path(panel.manifest['run_directory']).glob('提取失败文件*'))
        assert str(panel.failed_folder_button['state']) == 'normal'
        with patch('failed_files.os.startfile') as opened:
            panel.failed_folder_button.invoke(); opened.assert_called_once_with(str(bad.parent.resolve()))
        with patch('failed_files.messagebox.askyesno', return_value=True) as ask:
            start_without_prompt(panel); wait_for(root, panel); ask.assert_called_once()
        copied, = panel.manifest['failed_file_copies']['copied']
        assert Path(copied['copy']).read_bytes() == bad.read_bytes() == b'\xff\xfe\x00'
        with patch('failed_files.os.startfile') as opened:
            panel.failed_folder_button.invoke(); opened.assert_called_once_with(str(Path(copied['copy']).parent.resolve()))
        assert good.read_text(encoding='utf-8') == 'x,y\n1,2\n3,4'
        feedback_checks = {'text_failure_rate': 50, 'decline_creates_no_copy': True,
                           'accept_copies_failed_only': True, 'originals_preserved': True,
                           'failed_folder_button_verified': True}
    assert not callback_errors, callback_errors
    report = {'ok': True, 'version': root.title(), 'drop_checks': drop_checks, 'policy_checks': policy_checks,
              'frontend_checks': {'sidebar_navigation':True, 'inline_extra_fields':True, 'tk_callback_errors':callback_errors},
              'tabs': [app.notebook.tab(t, 'text') for t in app.notebook.tabs()],
              'manifest': manifest, 'all_columns_export': str(target), 'export_checks': checks, 'text_manifest': text_manifest,
              'data_preview_removed': True, 'secondary_actions_in_menu': True,
              'startup_queues_empty': True, 'chosen_output_used_without_prompt': True,
              'feedback_checks': feedback_checks}
    json_write(args.out / 'gui_smoke.json', report)
    return 0
