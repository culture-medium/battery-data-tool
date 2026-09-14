import copy
import json
from pathlib import Path
import tempfile
import threading
import unittest
from unittest.mock import patch

from openpyxl import load_workbook

from neware_batch import collect_files, run_batch
from neware_excel import export_workbook, worksheet_name
from neware_extract import HEADERS


def fake_result(path, **_kwargs):
    return {"source_file": str(path), "source_sha256": "testhash", "mass_mg": 1.0,
            "sample_name": {"value": "同名_样品", "source": "full_filename_stem_unconfirmed"},
            "audit": {"record_count": 2, "cycle_count": 1, "cycle_first_direction": "discharge"},
            "cycles": [{"cycle": 1, "charge_specific_capacity_mAh_g": 560.125,
                        "discharge_specific_capacity_mAh_g": 600.0, "efficiency_percent": 93.354,
                        "charge_energy_Wh": 0.001, "discharge_energy_Wh": 0.002,
                        "capacity_retention_percent": 100.0, "has_both_directions": True}]}


class BatchTests(unittest.TestCase):
    def test_default_excel_contains_four_numeric_columns(self):
        with tempfile.TemporaryDirectory() as temporary, patch("neware_batch.extract", side_effect=fake_result):
            root = Path(temporary)
            manifest = run_batch([root / "a.ndax"], root / "out")
            entry = manifest["entries"][0]
            workbook = load_workbook(manifest["workbook"])
            sheet = workbook[entry["sheet_name"]]
            self.assertEqual(list(sheet.values), [tuple(HEADERS[:4]), (1, 560.13, 600.0, 93.35)])
            self.assertTrue(all(cell.data_type == "n" for cell in sheet[2]))
            self.assertEqual([cell.number_format for cell in sheet[2]], ["0", "0.00", "0.00", "0.00"])
            self.assertEqual(sheet.freeze_panes, "B2")
            self.assertEqual(sheet.auto_filter.ref, "A1:D2")
            self.assertFalse(sheet.merged_cells)
            workbook.close()
            complete = json.loads(Path(entry["data"]).read_text(encoding="utf-8"))
            self.assertEqual(complete["cycles"][0]["discharge_energy_Wh"], 0.002)
            self.assertEqual(manifest["selected_columns"], [0, 1, 2, 3])

    def test_custom_columns_keep_values_aligned_and_source_unchanged(self):
        result = fake_result(Path("a.ndax"))
        before = copy.deepcopy(result)
        with tempfile.TemporaryDirectory() as temporary, patch("neware_batch.extract", return_value=result):
            root = Path(temporary)
            manifest = run_batch([root / "a.ndax"], root / "out", columns=[6, 1, 5])
            workbook = load_workbook(manifest["workbook"])
            sheet = workbook.active
            self.assertEqual(list(sheet.values), [tuple(HEADERS[i] for i in (1, 5, 6)), (560.13, 0.002, 100.0)])
            self.assertEqual([cell.number_format for cell in sheet[2]], ["0.00", "0.000000", "0.00"])
            workbook.close()
            # Re-export all cached samples with all columns; never invoke extract.
            with patch("neware_batch.extract", side_effect=AssertionError("Must use cached data")):
                path = root / "all.xlsx"
                export_workbook(manifest["entries"], path, range(7))
            workbook = load_workbook(path)
            self.assertEqual(workbook.active.max_column, 7)
            workbook.close()
        self.assertEqual(result, before)

    def test_empty_selection_does_not_start_or_write_results(self):
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            with self.assertRaises(ValueError):
                run_batch([root / "a.ndax"], root / "out", columns=[])
            self.assertFalse((root / "out").exists())

    def test_distinct_tests_with_same_name_preserve_both_results(self):
        with tempfile.TemporaryDirectory() as temporary, patch("neware_batch.extract", side_effect=fake_result):
            root = Path(temporary)
            manifest = run_batch([root / "a.ndax", root / "b.ndax"], root / "out")
            self.assertEqual(manifest["success"], 2)
            first, second = manifest["entries"]
            self.assertNotEqual(first["directory"], second["directory"])
            self.assertTrue(Path(first["data"]).exists())
            self.assertTrue(Path(second["data"]).exists())
            workbook = load_workbook(manifest["workbook"])
            self.assertEqual(workbook.sheetnames, ["同名_样品", "同名_样品 (2)"])
            self.assertEqual(len(list(Path(manifest["run_directory"]).rglob("*.xlsx"))), 1)
            self.assertEqual(len(manifest["excel_sheets"]), 2)
            workbook.close()

    def test_one_bad_file_does_not_abort_remaining_files(self):
        def extraction(path, **kwargs):
            if path.name == "bad.ndax":
                raise ValueError("Unsupported layout")
            return fake_result(path)
        with tempfile.TemporaryDirectory() as temporary, patch("neware_batch.extract", side_effect=extraction):
            root = Path(temporary)
            manifest = run_batch([root / "bad.ndax", root / "good.ndax"], root / "out")
            self.assertEqual((manifest["failed"], manifest["success"]), (1, 1))
            saved = json.loads((Path(manifest["run_directory"]) / "批次记录.json").read_text(encoding="utf-8"))
            self.assertIn("Unsupported layout", saved["entries"][0]["error"])
            workbook = load_workbook(saved["workbook"])
            self.assertEqual(workbook.sheetnames, [saved["entries"][1]["sheet_name"]])
            workbook.close()

    def test_stop_finishes_current_file_and_skips_remaining(self):
        cancel = threading.Event()
        def extraction(path, **kwargs):
            cancel.set()
            return fake_result(path)
        with tempfile.TemporaryDirectory() as temporary, patch("neware_batch.extract", side_effect=extraction):
            root = Path(temporary)
            manifest = run_batch([root / "a.ndax", root / "b.ndax"], root / "out", cancel=cancel)
            self.assertEqual(manifest["success"], 1)
            self.assertEqual(manifest["skipped"], 1)
            self.assertTrue(manifest["cancelled"])
            self.assertTrue(Path(manifest["workbook"]).exists())

    def test_no_empty_workbook_when_all_files_fail(self):
        with tempfile.TemporaryDirectory() as temporary, patch("neware_batch.extract", side_effect=ValueError("bad")):
            root = Path(temporary)
            manifest = run_batch([root / "bad.ndax"], root / "out")
            self.assertEqual(manifest["failed"], 1)
            self.assertIsNone(manifest["workbook"])
            self.assertEqual(list(root.rglob("*.xlsx")), [])

    def test_excel_save_failure_keeps_extraction_for_retry(self):
        with tempfile.TemporaryDirectory() as temporary, patch("neware_batch.extract", side_effect=fake_result):
            root = Path(temporary)
            with patch("neware_batch.export_workbook", side_effect=PermissionError("locked")):
                manifest = run_batch([root / "a.ndax"], root / "out")
            self.assertEqual(manifest["success"], 1)
            self.assertIsNone(manifest["workbook"])
            self.assertIn("locked", manifest["excel_error"])
            target = root / "retry.xlsx"
            export_workbook(manifest["entries"], target)
            original = target.read_bytes()
            with patch("neware_excel.os.replace", side_effect=PermissionError("locked")):
                with self.assertRaises(PermissionError):
                    export_workbook(manifest["entries"], target, range(7))
            self.assertEqual(target.read_bytes(), original)
            self.assertEqual(list(root.glob(".neware_*")), [])

    def test_sheet_names_preserve_arbitrary_names_and_resolve_excel_limits(self):
        used = set()
        self.assertEqual(worksheet_name("电池_多段-A_名字", used), "电池_多段-A_名字")
        self.assertEqual(worksheet_name("CS55-3.18-1C", used), "CS55-3.18-1C")
        self.assertEqual(worksheet_name("cs55-3.18-1c", used), "cs55-3.18-1c (2)")
        self.assertEqual(worksheet_name("'a[b]:c/d\\e?f*g'", used), "a_b__c_d_e_f_g")
        self.assertEqual(worksheet_name("History", used), "History_")
        self.assertEqual(worksheet_name("''", used), "测试数据")
        names = [worksheet_name("长名称_" * 10, used), worksheet_name("长名称_" * 10, used),
                 worksheet_name("🔋" * 30, used)]
        self.assertEqual(len(set(names)), 3)
        self.assertTrue(all(len(name.encode("utf-16-le")) <= 62 for name in names))

    def test_collect_mixed_case_extensions_nested_and_duplicate_paths(self):
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            sub = root / "子目录"
            sub.mkdir()
            top = root / "a.NDAX"
            nested = sub / "b.ndax"
            for path in [top, nested, root / "ignore.txt"]:
                path.write_text("fixture")
            self.assertEqual(set(collect_files([root, top])), {top, nested})
            self.assertEqual(collect_files([root], recursive=False), [top])


if __name__ == "__main__":
    unittest.main(verbosity=2)
