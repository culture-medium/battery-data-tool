"""Folder conversion using the same XLSX naming and styling conventions."""
from __future__ import annotations

from datetime import datetime
import json
import os
from pathlib import Path
import tempfile
import threading

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

from neware_batch import unique_directory, json_write
from neware_excel import worksheet_name
from text_extract import parse_text, find_peaks, validate_ranges

MERGE_MODES = ('全部文件合并', '按类型合并', '每个文件单独')


def collect_text_files(folder, recursive=True):
    if isinstance(folder, (list, tuple)):
        result, seen = [], set()
        for source in folder:
            source = Path(source)
            candidates = collect_text_files(source, recursive) if source.is_dir() else [source]
            for file in candidates:
                if file.is_file() and file.suffix.lower() in ('.txt', '.csv', '.tsv'):
                    full = file.resolve()
                    if str(full).casefold() not in seen:
                        result.append(full); seen.add(str(full).casefold())
        return result
    folder = Path(folder)
    if not folder.is_dir(): raise ValueError('请选择有效的输入文件夹。')
    return sorted(p.resolve() for p in (folder.rglob('*') if recursive else folder.glob('*'))
                  if p.is_file() and p.suffix.lower() in ('.txt', '.csv', '.tsv'))


def text_source_root(source, files):
    if isinstance(source, (str, Path)): return Path(source).resolve()
    try: return Path(os.path.commonpath([str(p.parent) for p in files])) if files else None
    except ValueError: return None  # Files on different drives retain full paths.


def text_identity(file, base):
    return file.relative_to(base).as_posix() if base is not None else file.as_posix()


def table_data(result, options):
    rows, headers = result['rows'], result['headers']
    if result['kind'] == 'EIS':
        if options.get('eis_simple'):
            return [["Z'/ohm", '-Z"/ohm']] + [[r[1], -r[2]] for r in rows], 1
        return [headers + ['-Z"/ohm']] + [r + [-r[2]] for r in rows], 1
    if result['kind'] == 'CV' and options.get('cv_split', True):
        cycles, n = result['cycles'], len(headers)
        first, second = [], []
        for i in range(len(cycles)):
            first += [f'第 {i+1} 圈'] + [None] * n
            second += headers + [None]
        output = [first[:-1], second[:-1]]
        for i in range(max(map(len, cycles))):
            row = []
            for cycle in cycles:
                row += (cycle[i] if i < len(cycle) else [None] * n) + [None]
            output.append(row[:-1])
        return output, 2
    return ([headers] if headers else []) + rows, int(headers is not None)


def add_sheet(workbook, name, rows, header_rows, used):
    if len(rows) > 1048576 or max(map(len, rows), default=0) > 16384:
        raise ValueError(f'{name} 超过 Excel 行列上限；请减少合并内容或关闭 CV 分圈。')
    sheet = workbook.create_sheet(worksheet_name(name, used))
    for row in rows:
        sheet.append(row)
    for row in sheet:
        for cell in row:
            if isinstance(cell.value, str):
                if len(cell.value) > 32767: raise ValueError('单个文本单元格超过 Excel 字符上限。')
                cell.data_type = 's'  # Source text beginning '=' remains literal.
            cell.font = Font(name='Microsoft YaHei', size=10, color='17324D')
            if cell.row <= header_rows:
                cell.font = Font(name='Microsoft YaHei', size=10, bold=True, color='FFFFFF')
                cell.fill = PatternFill('solid', fgColor='244568')
                cell.alignment = Alignment(wrap_text=True, vertical='center')
            elif isinstance(cell.value, float):
                cell.number_format = '0.##############'
    for index in range(1, sheet.max_column + 1):
        values = [sheet.cell(r, index).value for r in range(1, min(sheet.max_row, 30) + 1)]
        sheet.column_dimensions[get_column_letter(index)].width = 4 if not any(v is not None for v in values) else 23 if index > 1 else 30
    for index in range(1, header_rows + 1):
        sheet.row_dimensions[index].height = 72 if any(len(str(v or '')) > 30 for v in rows[index-1]) else 35
    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = f'A{header_rows+1}' if header_rows else None
    if header_rows == 1: sheet.auto_filter.ref = sheet.dimensions
    return sheet.title


def save_tables(tables, target):
    workbook = Workbook()
    workbook.remove(workbook.active)
    temporary = None
    try:
        used, mappings = set(), []
        for name, rows, headers in tables:
            mappings.append(add_sheet(workbook, name, rows, headers, used))
        target = Path(target)
        target.parent.mkdir(parents=True, exist_ok=True)
        with tempfile.NamedTemporaryFile(dir=target.parent, prefix='.text_', suffix='.xlsx', delete=False) as handle:
            temporary = Path(handle.name)
        workbook.save(temporary)
        os.replace(temporary, target)
        return mappings
    finally:
        workbook.close()
        if temporary: temporary.unlink(missing_ok=True)


def peak_tables(results, ranges):
    peaks = [(r['identity'], find_peaks(r['cycles'], ranges)) for r in results if r['kind'] == 'CV']
    max_cycles = max((len(p[0]['values']) for _, p in peaks), default=0)
    tables = []
    for log in (False, True):
        header = ['文件']
        for low, high, mode in ranges:
            for cycle in range(1, max_cycles+1):
                label = f'{low:g}～{high:g} V / {mode} / 第 {cycle} 圈'
                header.extend([label + (' / 电位(V)' if log else ' / 电流(A)'), label + (' / log10|I(A)|' if log else ' / 电位(V)')])
        rows = [header]
        for identity, intervals in peaks:
            row = [identity]
            for interval in intervals:
                for i in range(max_cycles):
                    value = interval['values'][i] if i < len(interval['values']) else None
                    row.extend(([value['potential'], value['logI']] if log else [value['current'], value['potential']]) if value else [None, None])
            rows.append(row)
        tables.append(('logI' if log else '电流+电位', rows, 1))
    return tables


def run_text_batch(folder, output, options=None, *, cancel=None, on_event=None):
    options = dict(options or {})
    ranges = validate_ranges(options.get('peak_ranges', []))
    merge = options.get('merge', MERGE_MODES[1])
    if merge not in MERGE_MODES: raise ValueError('无效的合并方式。')
    kinds = options.get('kinds', ['EIS', 'CV', '通用'])
    if not kinds: raise ValueError('请至少选择一种文本类型。')
    files = collect_text_files(folder, options.get('recursive', True))
    if not files: raise ValueError('文件夹中没有 TXT、CSV 或 TSV 文件。')
    base = text_source_root(folder, files)
    cancel, emit = cancel or threading.Event(), on_event or (lambda _: None)
    run = unique_directory(Path(output).resolve(), datetime.now().strftime('文本转换_%Y%m%d_%H%M%S'))
    manifest = {'run_directory': str(run), 'options': options, 'entries': [], 'workbooks': [], 'failed': 0,
                'success': 0, 'filtered': 0, 'cancelled': False, 'excel_errors': []}
    emit({'type': 'started', 'total': len(files), 'directory': str(run)})
    results, groups = [], {}
    for index, file in enumerate(files):
        if cancel.is_set():
            manifest['cancelled'] = True
            break
        item = {'index': index, 'source': str(file), 'identity': text_identity(file, base)}
        emit({'type': 'reading', **item})
        try:
            result = parse_text(file, delimiter=options.get('delimiter', '自动'), header=options.get('header', '自动'))
            item['kind'] = result['kind']
            if result['kind'] not in kinds:
                item['status'] = 'filtered'
                manifest['filtered'] += 1
            else:
                result['identity'] = item['identity']
                cache = run / '原始提取记录' / f'{index+1:04d}.json'
                cache.parent.mkdir(exist_ok=True)
                # Store the rows once; cycle boundaries reconstruct the split.
                stored = {k: v for k, v in result.items() if k != 'cycles'}
                stored['cycle_lengths'] = [len(c) for c in result['cycles']]
                json_write(cache, stored)
                table, header_rows = table_data(result, options)
                group = '文本数据' if merge == MERGE_MODES[0] else result['kind'] + '数据' if merge == MERGE_MODES[1] else f'{index+1:04d}_{file.stem}'
                groups.setdefault(group, []).append((result['name'], table, header_rows, item))
                results.append(result)
                item.update(status='success', rows=len(result['rows']), cycles=len(result['cycles']), data=str(cache), warnings=result['warnings'])
                manifest['success'] += 1
        except Exception as exc:
            item.update(status='error', error=str(exc))
            manifest['failed'] += 1
        manifest['entries'].append(item)
        emit({'type': 'file_done', **item})
    emit({'type': 'saving'})
    for name, group in groups.items():
        try:
            # Prefix guarantees unique filenames even when source stems coincide.
            target = run / (name[:110] + '.xlsx')
            mappings = save_tables([(n, rows, h) for n, rows, h, _ in group], target)
            manifest['workbooks'].append(str(target))
            for (_, _, _, item), sheet in zip(group, mappings): item.update(workbook=str(target), sheet_name=sheet)
        except Exception as exc:
            manifest['excel_errors'].append(f'{name}：{exc}')
    if ranges and any(r['kind'] == 'CV' for r in results):
        try:
            target = run / '峰分析.xlsx'
            save_tables(peak_tables(results, ranges), target)
            manifest['workbooks'].append(str(target))
        except Exception as exc:
            manifest['excel_errors'].append(f'峰分析：{exc}')
    manifest['unprocessed'] = len(files) - len(manifest['entries'])
    json_write(run / '批次记录.json', manifest)
    emit({'type': 'done', 'manifest': manifest})
    return manifest
