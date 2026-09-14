"""Offline XLSX export: one batch, one workbook, one worksheet per test."""
from __future__ import annotations

import json
import os
from pathlib import Path
import re
import tempfile

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from neware_columns import selected_columns
from battery_schema import profile, available_columns, display_values


def _shorten(text: str, limit: int) -> str:
    # Excel's limit is measured in UTF-16 code units; do not split an emoji.
    return text.encode("utf-16-le")[:limit * 2].decode("utf-16-le", errors="ignore")


def worksheet_name(name: str, used: set[str]) -> str:
    base = re.sub(r"[\[\]:*?/\\\x00-\x1f]", "_", str(name)).strip().strip("'") or "测试数据"
    if base.casefold() == "history":
        base += "_"
    candidate = _shorten(base, 31).rstrip("'")
    counter = 2
    while candidate.casefold() in used:
        suffix = f" ({counter})"
        candidate = _shorten(base, 31 - len(suffix)).rstrip("'") + suffix
        counter += 1
    used.add(candidate.casefold())
    return candidate


def export_workbook(entries: list[dict], target: Path, columns=None) -> list[dict]:
    """Read cached extraction data; write numeric columns without re-parsing NDAX.

    Save through a sibling temporary file. A failed save leaves any existing
    workbook intact, including when it is open and locked by Excel.
    """
    columns = selected_columns(columns)
    successful = [entry for entry in entries if entry["status"] == "success"]
    if not successful:
        raise ValueError("没有成功提取的文件，无法生成 Excel。")
    target = Path(target)
    if target.suffix.lower() != ".xlsx":
        raise ValueError("请使用 .xlsx 扩展名保存 Excel。")
    workbook = Workbook()
    workbook.remove(workbook.active)
    mappings, used = [], set()
    temporary = None
    font = Font(name="Microsoft YaHei", size=10, color="172B4D")
    header_font = Font(name="Microsoft YaHei", size=10, bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="244568")
    header_border = Border(right=Side(style="thin", color="FFFFFF"))
    try:
        for entry in successful:
            result = json.loads(Path(entry["data"]).read_text(encoding="utf-8"))
            if result.get('format') == 'land_cex' and not result.get('audit', {}).get('statistics_policy'):
                raise ValueError(f"{entry['sample_name']} 是旧版蓝电提取记录，请用原始 CEX 重新提取以校验循环口径。")
            schema = profile(result)
            effective = available_columns(columns, schema)
            if not effective:
                raise ValueError(f"{entry['sample_name']} 没有所选字段；请至少勾选一项通用字段。")
            rows = result["cycles"]
            if len(rows) > 1048575:
                raise ValueError(f"{entry['sample_name']} 超过 Excel 单个工作表的行数上限。")
            name = worksheet_name(entry["sample_name"], used)
            sheet = workbook.create_sheet(name)
            sheet.append([schema['headers'][i] for i in effective])
            for row in rows:
                rounded = display_values(row, schema)
                # Numeric values equal the verified display precision. Percent
                # columns contain e.g. 81.09 (header says %), never text or 0.8109.
                sheet.append([None if rounded[i] == '' else int(rounded[i]) if i == 0 else float(rounded[i]) for i in effective])
            for cell in sheet[1]:
                cell.font = header_font
                cell.fill = header_fill
                cell.border = header_border
                cell.alignment = Alignment(horizontal="center", vertical="center")
            sheet.row_dimensions[1].height = 27
            for j, source_column in enumerate(effective, 1):
                sheet.column_dimensions[get_column_letter(j)].width = 12 if source_column == 0 else 28
                precision = schema['precision'][source_column]
                number_format = "0" if precision == 0 else "0." + "0" * precision
                for cells in sheet.iter_rows(min_row=2, max_row=len(rows) + 1, min_col=j, max_col=j):
                    cell = cells[0]
                    cell.number_format = number_format
                    cell.font = font
                    cell.alignment = Alignment(horizontal="right", vertical="center")
            sheet.sheet_format.defaultRowHeight = 19
            sheet.sheet_view.showGridLines = False
            sheet.freeze_panes = "B2" if effective[0] == 0 and len(effective) > 1 else "A2"
            sheet.auto_filter.ref = sheet.dimensions
            mappings.append({"index": entry["index"], "sample_name": entry["sample_name"],
                             "sheet_name": name, "source": entry["source"], "rows": len(rows),
                             "columns": list(effective), "format": result.get('format', 'neware_ndax')})
        target.parent.mkdir(parents=True, exist_ok=True)
        with tempfile.NamedTemporaryFile(dir=target.parent, prefix=".neware_", suffix=".xlsx", delete=False) as handle:
            temporary = Path(handle.name)
        workbook.save(temporary)
        os.replace(temporary, target)
        return mappings
    finally:
        workbook.close()
        if temporary is not None:
            temporary.unlink(missing_ok=True)
